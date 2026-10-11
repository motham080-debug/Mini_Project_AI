import streamlit as st
import requests
import json
import io
import sqlite3
import hashlib
import re
import subprocess
import os
import tempfile
import imageio_ffmpeg
import fitz  # PyMuPDF สำหรับสกัดข้อความจาก PDF รูปภาพ
import speech_recognition as sr  # Google Web Speech API
from openai import OpenAI
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

# ==========================================
# 1. การตั้งค่า API Keys
# ==========================================
try:
    AIFORTHAI_API_KEY = st.secrets.get("AIFORTHAI_API_KEY", "PhC1TMiWfddbewzDGUtCZCgXoCg94uqZ")
    THAILLM_API_KEY = st.secrets.get("THAILLM_API_KEY", "gYkR4ZwL4Aq3gnkxHy9vQBv6owlXatza")
except Exception:
    AIFORTHAI_API_KEY = "PhC1TMiWfddbewzDGUtCZCgXoCg94uqZ"
    THAILLM_API_KEY = "gYkR4ZwL4Aq3gnkxHy9vQBv6owlXatza"



thaillm_client = OpenAI(
    api_key=THAILLM_API_KEY,
    base_url="https://api.thaillm.or.th/v1"
)

# ==========================================
# 2. ระบบฐานข้อมูลผู้ใช้ (SQLite Auth)
# ==========================================
@st.cache_resource
def init_db():
    with sqlite3.connect("users.db") as conn:
        c = conn.cursor()
        c.execute('''
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY,
                password TEXT NOT NULL,
                name TEXT NOT NULL
            )
        ''')
        conn.commit()

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def register_user(username, password, name):
    try:
        with sqlite3.connect("users.db") as conn:
            c = conn.cursor()
            c.execute("INSERT INTO users VALUES (?, ?, ?)", (username, hash_password(password), name))
            conn.commit()
            return True
    except sqlite3.IntegrityError:
        return False

def authenticate_user(username, password):
    with sqlite3.connect("users.db") as conn:
        c = conn.cursor()
        c.execute("SELECT name FROM users WHERE username = ? AND password = ?", (username, hash_password(password)))
        result = c.fetchone()
        return result[0] if result else None

init_db()

# ==========================================
# 3. Helper Functions (ประมวลผลไฟล์ & AI)
# ==========================================
def process_ocr(image_bytes):
    url = "https://api.aiforthai.in.th/ocr"
    headers = {"Apikey": AIFORTHAI_API_KEY}
    files = {"uploadfile": ("image.jpg", image_bytes, "image/jpeg")}
    try:
        res = requests.post(url, headers=headers, files=files, timeout=20)
        if res.status_code == 200:
            return res.json().get("text", "")
        return ""
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการเชื่อมต่อ OCR: {e}")
        return ""

def process_pdf(pdf_bytes):
    try:
        pdf_file = io.BytesIO(pdf_bytes)
        reader = PdfReader(pdf_file)
        text = ""
        for page in reader.pages:
            t = page.extract_text()
            if t: 
                text += t + "\n"
        
        if not text.strip():
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            ocr_texts = []
            
            progress_bar = st.progress(0, text="กำลังสกัดข้อความจากภาพใน PDF...")
            total_pages = len(doc)
            
            for page_index in range(total_pages):
                page = doc[page_index]
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("jpeg")
                
                page_text = process_ocr(img_bytes)
                if page_text.strip():
                    ocr_texts.append(f"--- หน้า {page_index + 1} ---\n" + page_text.strip())
                
                progress_bar.progress((page_index + 1) / total_pages)
            
            progress_bar.empty()
            text = "\n\n".join(ocr_texts)
            
        return text
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการอ่านไฟล์ PDF: {e}")
        return ""

def process_docx_input(docx_bytes):
    try:
        docx_file = io.BytesIO(docx_bytes)
        doc = Document(docx_file)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการอ่านไฟล์ Word: {e}")
        return ""

def process_stt(audio_bytes, file_name="audio.wav"):
    # 1. แปลงไฟล์เสียงเป็น WAV 16kHz Mono โดยใช้ imageio_ffmpeg
    send_bytes = audio_bytes
    try:
        ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
        
        file_ext = os.path.splitext(file_name)[1] or ".mp3"
        with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as in_file:
            in_file.write(audio_bytes)
            in_filename = in_file.name

        out_filename = in_filename + "_converted.wav"

        cmd = [ffmpeg_exe, "-y", "-i", in_filename, "-ar", "16000", "-ac", "1", out_filename]
        subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)

        with open(out_filename, "rb") as f:
            send_bytes = f.read()

        if os.path.exists(in_filename):
            os.remove(in_filename)
        if os.path.exists(out_filename):
            os.remove(out_filename)

    except Exception as e:
        st.error(f"ไม่สามารถแปลงไฟล์เสียงได้: {e}")
        return ""

    # 2. ถอดข้อความภาษาไทยด้วย Google Web Speech API (ฟรี ไม่ต้องใช้ API Key)
    recognizer = sr.Recognizer()
    try:
        audio_file = io.BytesIO(send_bytes)
        with sr.AudioFile(audio_file) as source:
            audio_data = recognizer.record(source)
            text = recognizer.recognize_google(audio_data, language="th-TH")
            return text
    except sr.UnknownValueError:
        st.warning("Google ไม่สามารถจับใจความเสียงนี้ได้ (เสียงอาจเบา หรือไม่มีเสียงพูด)")
        return ""
    except sr.RequestError as e:
        st.error(f"ไม่สามารถเชื่อมต่อเซิร์ฟเวอร์ Google ได้: {e}")
        return ""
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการถอดเสียง: {e}")
        return ""

def summarize_text(text):
    prompt = """คุณคือผู้ช่วยสรุปบทเรียนอัจฉริยะ ให้สรุปเนื้อหาต่อไปนี้อย่างเป็นระบบ เข้าใจง่าย:
- สรุปหัวข้อหลักและประเด็นสำคัญ (Key Points)
- จัดหมวดหมู่ด้วยหัวข้อ (# หรือ ##) และใช้ Bullet points (-)
- เน้นคำสำคัญด้วยตัวหนา (**ข้อความ**)
- ภาษาที่ใช้เข้าใจง่าย เหมาะสำหรับทบทวนก่อนสอบ"""
    try:
        res = thaillm_client.chat.completions.create(
            model="OpenThaiGPT-ThaiLLM-8B-Instruct-v7.2",
            messages=[{"role": "system", "content": prompt}, {"role": "user", "content": text}],
            temperature=0.3
        )
        return res.choices[0].message.content
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการสรุปเนื้อหา: {e}")
        return ""

def generate_quiz(text):
    prompt = """
    จากเนื้อหาที่กำหนด ให้สร้างข้อสอบปรนัย 3 ข้อ ในรูปแบบ JSON เพียวๆ ดังนี้:
    {
      "quizzes": [
        {
          "id": 1,
          "question": "คำถาม...",
          "options": ["ก. ตัวเลือก1", "ข. ตัวเลือก2", "ค. ตัวเลือก3", "ง. ตัวเลือก4"],
          "answer": "ก. ตัวเลือก1",
          "explanation": "เหตุผลเฉลย..."
        }
      ]
    }
    """
    try:
        res = thaillm_client.chat.completions.create(
            model="OpenThaiGPT-ThaiLLM-8B-Instruct-v7.2",
            messages=[{"role": "system", "content": prompt}, {"role": "user", "content": text}],
            temperature=0.4,
            response_format={"type": "json_object"}
        )
        content = res.choices[0].message.content.strip()
        content_clean = re.sub(r'^```json\s*|\s*```$', '', content, flags=re.MULTILINE)
        return json.loads(content_clean)
    except Exception as e:
        st.error(f"เกิดข้อผิดพลาดในการสร้างแบบทดสอบ: {e}")
        return {"quizzes": []}

def add_formatted_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            p.add_run(part)
    return p

def create_docx(text_content, user_name):
    doc = Document()
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"🎓 สรุปเนื้อหาการเรียน - คุณ{user_name}")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    lines = text_content.split('\n')
    for line in lines:
        line_str = line.strip()
        if not line_str: continue
        if line_str.startswith('# '):
            h = doc.add_heading(level=1)
            h.add_run(line_str[2:]).font.color.rgb = RGBColor(30, 58, 138)
        elif line_str.startswith('## '):
            h = doc.add_heading(level=2)
            h.add_run(line_str[3:]).font.color.rgb = RGBColor(30, 58, 138)
        elif line_str.startswith('- ') or line_str.startswith('* '):
            add_formatted_paragraph(doc, line_str[2:], style='List Bullet')
        elif re.match(r'^\d+\.\s', line_str):
            clean_text = re.sub(r'^\d+\.\s', '', line_str)
            add_formatted_paragraph(doc, clean_text, style='List Number')
        else:
            add_formatted_paragraph(doc, line_str)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==========================================
# 4. หน้าตา Web Application (Streamlit UI)
# ==========================================
st.set_page_config(page_title="Smart Study Companion Pro", page_icon="🎓", layout="wide")

if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False
if "user_name" not in st.session_state:
    st.session_state["user_name"] = ""

# --- หน้า Login / Register ---
if not st.session_state["authenticated"]:
    st.title("🎓 Smart Study Companion")
    st.caption("ระบบย่อยเนื้อหาการเรียนอัจฉริยะ สำหรับนักเรียน/นักศึกษา")
    
    auth_tab1, auth_tab2 = st.tabs(["🔑 เข้าสู่ระบบ (Login)", "📝 สมัครสมาชิก (Register)"])
    
    with auth_tab1:
        st.subheader("เข้าสู่ระบบเพื่อใช้งาน")
        login_user = st.text_input("ชื่อผู้ใช้ (Username)", key="login_user")
        login_pass = st.text_input("รหัสผ่าน (Password)", type="password", key="login_pass")
        if st.button("🚀 เข้าสู่ระบบ", use_container_width=True):
            user_fullname = authenticate_user(login_user, login_pass)
            if user_fullname:
                st.session_state["authenticated"] = True
                st.session_state["user_name"] = user_fullname
                st.success(f"ยินดีต้อนรับคุณ {user_fullname}!")
                st.rerun()
            else:
                st.error("ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง")

    with auth_tab2:
        st.subheader("ลงทะเบียนบัญชีใหม่")
        reg_name = st.text_input("ชื่อ-นามสกุล", key="reg_name")
        reg_user = st.text_input("ชื่อผู้ใช้ที่ต้องการ (Username)", key="reg_user")
        reg_pass = st.text_input("รหัสผ่าน (Password)", type="password", key="reg_pass")
        reg_pass_confirm = st.text_input("ยืนยันรหัสผ่าน", type="password", key="reg_pass_confirm")
        if st.button("✨ สมัครสมาชิก", use_container_width=True):
            if not reg_name or not reg_user or not reg_pass:
                st.warning("กรุณากรอกข้อมูลให้ครบทุกช่อง")
            elif reg_pass != reg_pass_confirm:
                st.error("รหัสผ่านทั้งสองช่องไม่ตรงกัน")
            else:
                if register_user(reg_user, reg_pass, reg_name):
                    st.success("สมัครสมาชิกสำเร็จ! กรุณาเข้าสู่ระบบที่แท็บ Login")
                else:
                    st.error("ชื่อผู้ใช้นี้มีผู้ใช้งานแล้ว")

# --- หน้าแอปพลิเคชันหลัก ---
else:
    st.sidebar.title("👤 บัญชีผู้ใช้")
    st.sidebar.write(f"สวัสดี, **คุณ{st.session_state['user_name']}**")
    if st.sidebar.button("🚪 ออกจากระบบ (Logout)", use_container_width=True):
        st.session_state["authenticated"] = False
        st.session_state["user_name"] = ""
        st.rerun()

    st.title("🎓 Smart Study Companion")
    
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📂 สแกน/อัปโหลดเอกสาร", 
        "🎙️ ถอดเสียงคลิปสอน (STT)", 
        "📝 สรุปเนื้อหา & ดาวน์โหลด", 
        "🎯 ทำแบบทดสอบ (Quiz)", 
        "💬 ถาม-ตอบกับ AI Tutor"
    ])

    with tab1:
        st.header("📂 อัปโหลดชีทเรียน / เอกสารการสอน")
        uploaded_file = st.file_uploader("เลือกไฟล์รูปภาพ (JPG, PNG), PDF หรือ Word (.docx)", type=["jpg", "png", "jpeg", "pdf", "docx"])
        if uploaded_file and st.button("🚀 สกัดข้อความจากเอกสาร"):
            ext = uploaded_file.name.split(".")[-1].lower()
            with st.spinner("กำลังอ่านข้อความ..."):
                if ext in ["jpg", "png", "jpeg"]:
                    extracted_text = process_ocr(uploaded_file.getvalue())
                elif ext == "pdf":
                    extracted_text = process_pdf(uploaded_file.getvalue())
                elif ext == "docx":
                    extracted_text = process_docx_input(uploaded_file.getvalue())
                else:
                    extracted_text = ""
            
            if extracted_text and extracted_text.strip():
                st.session_state["study_text"] = extracted_text
                st.success("ดึงข้อความสำเร็จ!")
                st.text_area("ข้อความที่ได้:", value=extracted_text, height=250)
            else:
                st.warning("ไม่สามารถสกัดข้อความจากเอกสารนี้ได้ หรือไม่พบข้อความในไฟล์")

    with tab2:
        st.header("🎙️ อัปโหลดไฟล์เสียงอัดการสอน")
        audio_file = st.file_uploader("เลือกไฟล์เสียง (.wav / .mp3)", type=["wav", "mp3"])
        
        if audio_file:
            st.audio(audio_file)
            
            if st.button("🚀 ถอดข้อความจากไฟล์เสียง"):
                with st.spinner("กำลังแปลงเสียงเป็นข้อความ..."):
                    transcribed_text = process_stt(audio_file.getvalue(), audio_file.name)
                
                if transcribed_text and transcribed_text.strip():
                    st.session_state["study_text"] = transcribed_text
                    st.success("ถอดเสียงสำเร็จ!")
                    st.text_area("ข้อความจากเสียง:", value=transcribed_text, height=200)
                else:
                    st.warning("ไม่สามารถถอดข้อความจากไฟล์เสียงนี้ได้")

    with tab3:
        st.header("📝 สรุปเนื้อหาบทเรียน & ดาวน์โหลด")
        raw_text = st.text_area("เนื้อหาบทเรียนต้นฉบับ:", value=st.session_state.get("study_text", ""), height=150)
        st.session_state["study_text"] = raw_text

        if st.button("✨ ให้ AI สรุปเนื้อหาบทเรียน ", use_container_width=True):
            if raw_text.strip():
                with st.spinner("ThaiLLM กำลังสรุปเนื้อหาบทเรียน..."):
                    summary_result = summarize_text(raw_text)
                    st.session_state["summary_text"] = summary_result
                    st.success("สรุปเนื้อหาสำเร็จ!")
            else:
                st.warning("กรุณาใส่เนื้อหาบทเรียนก่อนทำการสรุป")

        if "summary_text" in st.session_state and st.session_state["summary_text"]:
            st.write("---")
            st.subheader("📌 ผลการสรุปเนื้อหาโดย AI:")
            summary_display = st.text_area("แก้ไข/ตรวจสอบเนื้อหาสรุป:", value=st.session_state["summary_text"], height=250)
            st.session_state["summary_text"] = summary_display
            
            col1, col2 = st.columns(2)
            with col1:
                st.download_button("📄 ดาวน์โหลดสรุปเป็น Text (.txt)", summary_display, file_name="summary.txt", mime="text/plain", use_container_width=True)
            with col2:
                docx_file = create_docx(summary_display, st.session_state['user_name'])
                st.download_button("📝 ดาวน์โหลดสรุปเป็น Word (.docx)", docx_file, file_name="summary.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

    with tab4:
        st.header("🎯 แบบทดสอบประเมินความรู้")
        if st.button("🎲 สร้างข้อสอบใหม่จากเนื้อหา"):
            if st.session_state.get("study_text"):
                for k in list(st.session_state.keys()):
                    if k.startswith("q_"):
                        del st.session_state[k]
                        
                with st.spinner("ThaiLLM กำลังออกแบบข้อสอบ..."):
                    quiz_data = generate_quiz(st.session_state["study_text"])
                    st.session_state["quizzes"] = quiz_data.get("quizzes", [])
            else:
                st.warning("กรุณาเพิ่มเนื้อหาใน Tab 1 หรือ 2 ก่อน")

        if "quizzes" in st.session_state and st.session_state["quizzes"]:
            user_answers = {}
            with st.form("quiz_form"):
                for q in st.session_state["quizzes"]:
                    st.subheader(f"ข้อที่ {q['id']}: {q['question']}")
                    user_answers[q['id']] = st.radio("เลือกคำตอบ:", q['options'], key=f"q_{q['id']}")
                    st.write("---")
                
                submitted = st.form_submit_button("ส่งคำตอบและตรวจคะแนน")
                if submitted:
                    score = sum([1 for q in st.session_state["quizzes"] if user_answers.get(q['id']) == q['answer']])
                    for q in st.session_state["quizzes"]:
                        if user_answers.get(q['id']) == q['answer']:
                            st.success(f"ข้อ {q['id']}: ถูกต้อง! 🎉 {q['explanation']}")
                        else:
                            st.error(f"ข้อ {q['id']}: ผิด ❌ (เฉลย: {q['answer']}) {q['explanation']}")
                    st.info(f"🏆 คะแนนของคุณ: {score} / {len(st.session_state['quizzes'])} คะแนน")

    with tab5:
        st.header("💬 พูดคุยสอบถามประเด็นที่สงสัยกับ AI Tutor")
        if "chat_history" not in st.session_state:
            st.session_state["chat_history"] = []

        for msg in st.session_state["chat_history"]:
            with st.chat_message(msg["role"]):
                st.write(msg["content"])

        user_query = st.chat_input("พิมพ์คำถามที่สงสัยเกี่ยวกับบทเรียน...")
        if user_query:
            st.session_state["chat_history"].append({"role": "user", "content": user_query})
            with st.chat_message("user"):
                st.write(user_query)

            context = st.session_state.get("study_text", "ไม่มีบริบทเนื้อหา")
            system_prompt = f"คุณคือติวเตอร์ AI ตอบคำถามสั้นกระชับเข้าใจง่าย อิงตามบริบทเนื้อหานี้:\n{context}"
            with st.chat_message("assistant"):
                with st.spinner("กำลังคิดคำตอบ..."):
                    try:
                        res = thaillm_client.chat.completions.create(
                            model="OpenThaiGPT-ThaiLLM-8B-Instruct-v7.2",
                            messages=[{"role": "system", "content": system_prompt}, {"role": "user", "content": user_query}]
                        )
                        bot_reply = res.choices[0].message.content
                    except Exception as e:
                        bot_reply = f"เกิดข้อผิดพลาดในการตอบคำถาม: {e}"
                        
                    st.write(bot_reply)
                    st.session_state["chat_history"].append({"role": "assistant", "content": bot_reply})